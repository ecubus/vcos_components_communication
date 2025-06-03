# This Python file uses the following encoding: utf-8
# ##############################################################################
# Copyright (c) 2025 Li Auto Inc. and its affiliates
# Licensed under the Apache License, Version 2.0(the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
# ##############################################################################
import os
from basic_func_module.file_handler.excel_parser.import_excel_base import ImportExcelBase
from openpyxl.reader.excel import load_workbook
from func_module.health_monitor.logger import logger


class ComExcelParser(ImportExcelBase):
    def __init__(self) -> None:
        super().__init__()
        pass

    def execute_import_excel_command(self, excel_path_list) -> dict:
        dict = {}
        if not excel_path_list:
            return None
        for excel_path in excel_path_list:
            excel_name = os.path.basename(excel_path)
            if excel_name == 'ComCfg.xlsx':
                result = self.parse_excel_file(excel_path)
                if result is None or result["excel_dict"] is None:
                    logger.ui_error(f"{excel_path} load ComCfg.xlsx failed, result is None!")
                    return None
                if self.excel_validate(result["excel_dict"]) is False:
                    logger.ui_error(f"{excel_path} ComCfg.xlsx Validation Failed!")
                    return None
                dict.update({excel_name: result["result_dict"]})
        return dict

    def parse_excel_file(self, file_path) -> dict:
        com_excel_dict = {"excel_dict": {}, "result_dict": {}}
        workbook = load_workbook(file_path)
        sheet = workbook["Sheet1"]
        # 固定excel文件为4列
        if sheet.max_column != 4:
            return
        project_name = self.get_cell_value(sheet, 2, 1)
        for row in range(3, sheet.max_row + 1):
            dest_list = []
            row_value_list = self.get_row_value(sheet, row)
            dict = {
                "Name": row_value_list[0],
                "Type": row_value_list[1],
                "Source": row_value_list[2],
                "Target": row_value_list[3],
            }
            com_excel_dict["excel_dict"].update({row: dict})
            for dest in row_value_list[3].split(','):
                # RX / TX 怎么确认？
                dest_list.append(self.get_pdu_name(dict["Name"], project_name, dest) + "_TX")
            # if dict["Source"] not in dict["Target"]:
            #     dest_list.append(self.get_pdu_name(dict["Name"], project_name, dict["Source"]) + "_TX")
            src_name = self.get_pdu_name(dict["Name"], project_name, dict["Source"]) + "_RX"
            com_excel_dict["result_dict"].update({src_name: dest_list})

        return com_excel_dict

    def get_pdu_name(self, message_name, project_name, can_lin_name) -> str:
        if can_lin_name in ["BDCAN2", "BLECAN"]:
            can_lin_name += "_Ctrl"
        return message_name + "_o" + project_name + "_" + can_lin_name

    def excel_validate(self, dict: dict) -> bool:
        validate_result = True
        for key, value in dict.items():
            validate = True
            if value["Type"] != "PDU Gateway" and value["Type"] != "Signal Gateway":
                validate = False
            if "CAN" not in value["Source"] and "LIN" not in value["Source"]:
                validate = False
            target_list = value["Target"].split(',')
            for target in target_list:
                if "CAN" not in target and "LIN" not in target:
                    validate = False
            if validate is False:
                validate_result = False
                logger.info(f"{key, value} is invalid info!")
        return validate_result